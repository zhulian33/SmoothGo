from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("/Users/zhulian/Documents/ChatGPT/日常办公/SmoothGo/项目介绍.docx")

INK = "0B1F1A"
INK_SOFT = "2D4039"
MUTED = "687A72"
TEAL = "0F9D86"
TEAL_SOFT = "EAF8F3"
LINE = "DCE5DF"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"


def set_run(run, size=11, bold=False, color=INK_SOFT, italic=False):
    run.font.name = "Arial"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    rFonts.set(qn("w:cs"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style(style, size, color, bold=False, before=0, after=6, line=1.2):
    style.font.name = "Arial"
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.rFonts
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    rFonts.set(qn("w:cs"), "Arial Unicode MS")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), LINE)
        borders.append(element)
    tblPr.append(borders)


def add_page_number(paragraph):
    paragraph.add_run("\t")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, TEAL_SOFT)
    r = p.add_run(text)
    set_run(r, size=11, bold=True, color=TEAL)


def add_body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=10.8, color=INK_SOFT)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run(r, size=10.8, color=INK_SOFT)


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_overview_table(doc):
    rows = [
        ("项目名称", "SmoothGo（畅游通）"),
        ("项目类型", "旅游决策服务 + 产品原型"),
        ("试点地点", "长春，之后拓展到西安、重庆、长沙"),
        ("目标用户", "自由行游客、家庭、时间紧张的商务出差者"),
        ("核心差异", "不只推荐去哪，而是根据实时拥挤度告诉你现在去哪、要不要改"),
        ("商业模式", "代排佣金 + 商家推广 + 高级订阅"),
        ("8 周试点", "10 个合作场地、1,000 个用户、8 周"),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [Cm(3.6), Cm(12.8)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    header = table.rows[0].cells
    header[0].text = ""
    header[1].text = ""
    p = header[0].paragraphs[0]
    r = p.add_run("项目概览")
    set_run(r, size=9, bold=True, color=TEAL)
    p = header[1].paragraphs[0]
    r = p.add_run("一句话看懂 SmoothGo")
    set_run(r, size=9, bold=True, color=TEAL)
    for cell in header:
        shade_cell(cell, TEAL_SOFT)
    for i, (label, value) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text = ""
        cells[1].text = ""
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run(r, size=9.5, bold=True, color=INK)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run(r, size=10, color=INK_SOFT)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    set_style(doc.styles["Normal"], 10.8, INK_SOFT, after=6, line=1.2)
    set_style(doc.styles["Heading 1"], 14, BLUE, bold=True, before=14, after=6, line=1.1)
    set_style(doc.styles["Heading 2"], 12, BLUE_DARK, bold=True, before=10, after=4, line=1.1)
    set_style(doc.styles["List Bullet"], 10.8, INK_SOFT, after=4, line=1.18)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("SmoothGo（畅游通）项目介绍")
    set_run(r, size=22, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("商业计划项目 · 长春试点 · 2026 小学期 Final Presentation")
    set_run(r, size=10, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("长春师范大学 / Southern Cross University")
    set_run(r, size=8.5, color=MUTED)

    add_callout(
        doc,
        "用实时客流和排队数据，帮游客决定“现在去哪、要不要改路线”。",
    )

    add_overview_table(doc)
    doc.add_paragraph()

    add_heading(doc, "一、项目大概是什么")
    add_body(
        doc,
        "SmoothGo（畅游通）是一个面向中国城市游客的智能行程决策服务。"
        "用户输入兴趣、节奏、同行人、开始时间和游玩时长，系统结合实时拥挤度、排队时间、距离和营业信息，"
        "生成一条可执行的行程；当某个点太挤时，用户可以一键 Replan，换到更合适的地方。",
    )
    add_body(
        doc,
        "它不是另一个订票网站，也不是单纯的排队工具。它解决的是游客在旅途中不断做决定的问题："
        "现在去哪、怎么走、要不要改。App / 小程序只是产品原型和演示工具，"
        "整个汇报里只占 2 页；商业计划才是主体。",
    )

    add_heading(doc, "二、在哪里做")
    add_body(
        doc,
        "第一阶段放在长春。原因很现实：我们是本地学生，熟悉城市；用户获取成本低；"
        "景点和餐饮比较集中，适合先跑一个小试点；学校也有学生和兼职人员资源。",
    )
    add_bullets(
        doc,
        [
            "试点景点：净月潭、长影世纪城、伪满皇宫博物院等。",
            "试点餐饮/街区：这有山、桂林路等。",
            "8 周试点目标：10 个合作场地、1,000 个用户。",
            "后续扩展：西安、重庆、长沙。",
        ],
    )
    add_body(
        doc,
        "我们不会一上来就说“全国平台”。先把一个城市跑通，再复制到下一个城市，这样更可信。",
    )

    add_heading(doc, "三、定位是什么")
    add_body(
        doc,
        "一句话定位：面向自由行游客的“排队感知型行程决策服务”。",
    )
    add_body(doc, "目标用户分三类：")
    add_bullets(
        doc,
        [
            "自由行游客：想要灵活，但不想把时间浪费在临时决策上。",
            "家庭：需要更现实的节奏和更短的排队。",
            "商务出差者：时间短、时间价值高，愿意为效率付费。",
        ],
    )
    add_body(
        doc,
        "和现有平台的区别：携程强在订票和发现，美团/大众点评强在本地餐饮和评价，"
        "内容平台强在灵感，单场地排队工具只解决一个场地。SmoothGo 做的是跨场地、实时、动态改线的决策层。",
    )
    add_body(
        doc,
        "边界也要讲清楚：我们不卖票、不承诺优先入场；景区先做信息监测，"
        "代排只在明确允许的餐厅或非严格排队场地提供。",
    )

    doc.add_page_break()

    add_heading(doc, "四、为什么值得做")
    add_body(
        doc,
        "游客最宝贵的不是钱，而是时间。一次 45 分钟的排队，可能让一顿饭、一个博物馆或最后一班车全部乱掉。"
        "更麻烦的是，游客往往不知道“现在该不该去下一个点”。",
    )
    add_bullets(
        doc,
        [
            "痛点清楚：时间浪费在排队和临时决策上。",
            "现有产品错位：大部分 App 告诉你“有什么”，很少告诉你“现在该去哪”。",
            "模式轻：不建景区、不买设备、不压门票库存。",
            "可试点：一个城市、10 个场地、1,000 个用户就能验证。",
            "可扩展：跑通后可以复制到其他旅游城市。",
        ],
    )

    add_heading(doc, "五、具体怎么做")
    add_heading(doc, "1. 产品流程", level=2)
    add_bullets(
        doc,
        [
            "用户选择兴趣、节奏、同行人、开始时间和游玩时长。",
            "系统生成一条带交通、吃饭和排队时间的行程。",
            "页面显示每个地点的拥挤度和预计排队时间。",
            "某个点太挤时，用户点 Replan，系统替换成更合适的地点。",
            "在允许的场地，用户可以选择可选的 Queue Assist。",
        ],
    )

    add_heading(doc, "2. 数据来源", level=2)
    add_bullets(
        doc,
        [
            "场地合作数据：合作景区/餐厅愿意提供的实时信息。",
            "公开旅游数据：政府、平台和行业报告。",
            "用户主动上报：用户看到排队情况可以上报换积分。",
        ],
    )
    add_body(
        doc,
        "初期数据不可能很准，所以我们会显示“估算”和置信度，而不是假装每个数字都是实时的。",
    )

    add_heading(doc, "3. 运营方式", level=2)
    add_bullets(
        doc,
        [
            "场地：学校孵化器或联合办公空间，不租门店。",
            "供给：实名、评分、押金的兼职跑腿人员。",
            "交接：用户和跑腿双方扫码确认，平台结算。",
            "原则：不建景区、不买设备、不压票。",
        ],
    )

    add_heading(doc, "4. 怎么赚钱", level=2)
    add_bullets(
        doc,
        [
            "代排佣金：完成订单后抽 15—20%。",
            "商家推广：合作场地优先展示。",
            "高级订阅：更详细的预测和多日规划。",
            "促销：首次免费、淡季/工作日折扣。",
        ],
    )
    add_body(
        doc,
        "所有财务数字都是 illustrative assumptions（示意性假设），不是实际业绩。",
    )

    add_heading(doc, "5. 怎么获客", level=2)
    add_bullets(
        doc,
        [
            "短视频：拍“这个景点现在排多久”。",
            "场地二维码：游客在门口扫码就能用。",
            "旅行社/自由行增值服务：打包给自由行游客。",
            "校园大使：先从同学和本地年轻人开始。",
        ],
    )

    add_heading(doc, "6. 8 周试点", level=2)
    add_bullets(
        doc,
        [
            "10 个合作场地。",
            "1,000 个用户。",
            "观察 4 个指标：完成行程数、Replan 使用率、节省的排队分钟、重复使用率。",
            "试点后再决定是否扩展到西安、重庆、长沙。",
        ],
    )

    add_heading(doc, "六、风险与边界")
    add_bullets(
        doc,
        [
            "数据不准：先用用户上报和可见排队估算，标出置信度。",
            "代排公平性：先做餐厅和非严格排队，景区只做信息监测。",
            "人员管理：实名、评分、押金、扫码交接。",
            "淡旺季：增加本地生活场景和淡季优惠。",
        ],
    )
    add_body(
        doc,
        "这些风险不会让项目“不能做”，但它们决定了试点怎么做。"
        "先把边界讲清楚，比假装没有风险更可信。",
    )

    add_heading(doc, "七、一句话总结")
    add_body(
        doc,
        "SmoothGo 想做的不是又一个旅游 App，而是一个根据实时客流帮游客做决定的商业服务。"
        "先在长春证明它真的能省时间、有人愿意付费，再复制到更多城市。",
    )

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), alignment=2)
    r = p.add_run("SmoothGo 项目介绍")
    set_run(r, size=8, color=MUTED)
    add_page_number(p)
    for run in p.runs:
        set_run(run, size=8, color=MUTED)

    cp = doc.core_properties
    cp.title = "SmoothGo（畅游通）项目介绍"
    cp.subject = "商业计划项目介绍"
    cp.author = "SmoothGo student team"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
